r"""Construye un caso (JSON) desde un Excel de propuesta del SII y calcula.

Uso:
    venv\Scripts\python.exe construir_caso.py <ruta.xlsx> [salida.json]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from app.schemas.orquestador import SimuladorGlobalRequest  # noqa: E402
from app.services.orquestador import OrquestadorService  # noqa: E402


def _leer_pares(ws, col_id, col_val):
    out = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if len(fila) < max(col_id, col_val) + 1:
            continue
        ident = fila[col_id]
        if ident is None:
            continue
        out[str(ident).strip()] = fila[col_val]
    return out


def _filas_retiros(ws):
    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila[0] is None:
            continue
        filas.append({
            "rut": str(fila[0]).strip(),
            "usufructuario": fila[1],
            "acciones": fila[2],
            "f1_fecha": str(fila[3]) if fila[3] is not None else "",
            "f1_monto": fila[4],
            "f1_isfut_h": fila[5],
            "f1_isfut_a": fila[6],
            "saldo": fila[7],
            "f2_fecha": str(fila[8]) if fila[8] is not None else "",
            "f2_monto": fila[9],
            "f2_isfut_h": fila[10],
            "f2_isfut_a": fila[11],
            "es_registro_nuevo": True,
        })
    return filas


def main() -> None:
    ruta = Path(sys.argv[1])
    wb = openpyxl.load_workbook(ruta, data_only=True)
    vectores = _leer_pares(wb["Vectores"], 2, 4)
    externos = _leer_pares(wb["Calculadora"], 2, 3)
    retiros = _filas_retiros(wb["Retiros"]) if "Retiros" in wb.sheetnames else []
    if len(sys.argv) > 3 and sys.argv[3] == "sinretiros":
        retiros = []

    payload = {
        "at": "2026",
        "patrimonio_personal": True,
        "mostrar_formulas": False,
        "vectores": vectores,
        "externos": externos,
        "digitados": {"retiros": {"filas": retiros}},
    }
    req = SimuladorGlobalRequest.model_validate(payload)
    resp = OrquestadorService().calcular_simulacion(req)
    print("vectores:", len(vectores), "| externos:", len(externos), "| retiros:", len(retiros))
    print("total7 =", resp.ingresos.totales.fila_7_total)
    print("total8 =", resp.egresos.totales.fila_8_total)
    r = resp.rli
    print("RLI  9 =", r.v9, "| 9.1 =", r.v9_1, "| 9.2 =", r.v9_2, "| 9.21 =", r.v9_21)
    print("RLI  9.3 =", r.v9_3, "| max94 =", r.max94, "| 9.4 =", r.v9_4, "| 9.5 =", r.v9_5, "| 9.6 =", r.v9_6)
    print("condicion_ok =", r.condicion_ok, "| cuadro =", r.cuadro)
    for f in resp.ingresos.filas:
        if f.monto_ingreso_percibido not in (None, 0) and f.codigo in ("7.14", "7.19", "7.20", "7.1"):
            print("  ing", f.codigo, "B=", f.ingresos_ano, "F=", f.monto_ingreso_percibido)

    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print("caso guardado en", sys.argv[2])


if __name__ == "__main__":
    main()

