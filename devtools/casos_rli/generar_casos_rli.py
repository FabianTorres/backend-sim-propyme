r"""Generador de casos Excel para certificar la pantalla RLI (Pagina 4).

Crea archivos .xlsx con el patron del Excel de propuesta del SII (hojas
Vectores / Calculadora / Retiros) para ejercitar los cuadros de la
Determinacion RLI en el frontend. Cada caso se valida contra el motor
(OrquestadorService) y la hoja Guia documenta los resultados esperados.

Uso:
    venv\Scripts\python.exe devtools\casos_rli\generar_casos_rli.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ_BACKEND = Path(__file__).resolve().parents[2]
if str(RAIZ_BACKEND) not in sys.path:
    sys.path.insert(0, str(RAIZ_BACKEND))

from app.schemas.orquestador import SimuladorGlobalRequest  # noqa: E402
from app.services.orquestador import OrquestadorService  # noqa: E402

RAIZ_REPO = Path(__file__).resolve().parents[3]
SALIDA = RAIZ_REPO / "simulador-propyme-ui" / "docs" / "casos_rli"

RUT_CASO = "70000001-0"

GLOSAS = {
    "Vx010934": "Ingresos 7.18 Intereses indirectos percibidos",
    "Vx012214": "Egresos 8.4 Compras y/o servicios internos del giro",
    "Vx013649": "Egresos 8.26 Partidas art.21 inc.1 pagadas",
    "Vx010213": "Tope base imponible (condicion incentivo 14E)",
    "Vx013013": "Base IDPC voluntario propuesto (tope 9.5)",
    "Vx011930": "Vector condicion incentivo al ahorro (14E)",
}

RETIROS_CASO5 = [
    {
        "rut": "1567340-0",
        "usufructuario": 0,
        "acciones": 0,
        "f1_fecha": "03032025",
        "f1_monto": 20000000,
        "f1_isfut_h": 5000000,
        "f1_isfut_a": 5000000,
        "saldo": 0,
        "f2_fecha": "",
        "f2_monto": 0,
        "f2_isfut_h": 0,
        "f2_isfut_a": 0,
    }
]

CASOS = [
    {
        "archivo": "RLI_caso1_t1_con_14E.xlsx",
        "titulo": "RLI 1 - Cuadro N1 con incentivo 14E (responder SI)",
        "descripcion": "9.1 > 0 y la condicion 14E se cumple. Al responder SI al modal se despliega el Cuadro N1 y el detalle 1.1.",
        "cuadro_esperado": "pendiente",
        "respuesta_ui": "SI",
        "vectores": {
            "Vx010934": 100000000,
            "Vx012214": 30000000,
            "Vx013649": 5000000,
            "Vx010213": 100000000,
            "Vx013013": 10000000,
        },
        "pasos": [
            "Responder SI en el modal 14E -> Cuadro N1.",
            "El detalle 1.1 muestra 9.3 y el maximo 9.4 (max94).",
        ],
    },
    {
        "archivo": "RLI_caso2_t3_condicion_falla.xlsx",
        "titulo": "RLI 2 - Cuadro N3 (condicion 14E no se cumple)",
        "descripcion": "9.1 > 0 pero la condicion 14E es 0: NO aparece el modal y se va directo al Cuadro N3.",
        "cuadro_esperado": "t3",
        "respuesta_ui": None,
        "vectores": {
            "Vx010934": 50000000,
            "Vx012214": 10000000,
            "Vx011930": 999000000,
            "Vx013013": 5000000,
        },
        "pasos": ["Sin modal 14E. Se despliega directamente el Cuadro N3."],
    },
    {
        "archivo": "RLI_caso3_t2_perdida.xlsx",
        "titulo": "RLI 3 - Cuadro N2 (perdida tributaria 9.2 > 0)",
        "descripcion": "La base (7 - 8 + 8.26) es negativa: 9.1 = 0 y 9.2 > 0. Se despliega el Cuadro N2 y no hay modal.",
        "cuadro_esperado": "t2",
        "respuesta_ui": None,
        "vectores": {
            "Vx010934": 10000000,
            "Vx012214": 40000000,
            "Vx010213": 100000000,
        },
        "pasos": ["Sin modal. Se despliega el Cuadro N2 (RLI negativa)."],
    },
    {
        "archivo": "RLI_caso4_pendiente_override_9_4.xlsx",
        "titulo": "RLI 4 - Modal pendiente y override del 9.4",
        "descripcion": "9.1 > 0, condicion 14E cumplida y Vx013013 = 0. Ejercita el estado pendiente y la edicion manual del 9.4 (override con cota).",
        "cuadro_esperado": "pendiente",
        "respuesta_ui": "SI",
        "vectores": {
            "Vx010934": 40000000,
            "Vx012214": 5000000,
            "Vx010213": 100000000,
        },
        "pasos": [
            "Estado inicial: modal 14E pendiente (cuadro provisorio).",
            "Responder SI -> 9.4 = 17.500.000 (maximo propuesto).",
            "Editar 9.4 a 8.000.000 -> 9.6 baja a 27.000.000 (dentro de cota).",
            "Editar 9.4 a 50.000.000 -> se recorta a 17.500.000.",
        ],
    },
    {
        "archivo": "RLI_caso5_retiros_reducen_9_3.xlsx",
        "titulo": "RLI 5 - Retiros (RET30) reducen el 9.3",
        "descripcion": "9.1 > 0 con retiros (RET30 = 20.000.000) que reducen el 9.3 invertido en la empresa y, con ello, el tope 9.4.",
        "cuadro_esperado": "pendiente",
        "respuesta_ui": "SI",
        "vectores": {
            "Vx010934": 100000000,
            "Vx012214": 30000000,
            "Vx010213": 100000000,
        },
        "retiros": RETIROS_CASO5,
        "pasos": [
            "La hoja Retiros carga el socio y su retiro (RET30 = 20.000.000).",
            "9.3 = 9.1 - RET30 - 8.26 = 50.000.000 (menor que 9.1).",
            "Responder SI -> 9.4 = max94 = 25.000.000; 9.6 = 45.000.000.",
        ],
    },
]


def _calcular(vectores, externos, retiros=None, acoge=None, deduccion=0):
    digitados = {"rli": {"acoge_14e": acoge, "deduccion_14e": deduccion}}
    if retiros:
        digitados["retiros"] = {"filas": retiros}
    payload = {
        "at": "2026",
        "mostrar_formulas": False,
        "vectores": vectores,
        "externos": externos or {},
        "digitados": digitados,
    }
    req = SimuladorGlobalRequest.model_validate(payload)
    return OrquestadorService().calcular_simulacion(req).rli


def _tabla(rli):
    return [
        ("9", str(rli.v9)),
        ("9.1", str(rli.v9_1)),
        ("9.2", str(rli.v9_2)),
        ("9.21", str(rli.v9_21)),
        ("9.3", str(rli.v9_3)),
        ("max94", str(rli.max94)),
        ("9.4", str(rli.v9_4)),
        ("9.5", str(rli.v9_5)),
        ("9.6", str(rli.v9_6)),
        ("condicion_ok", str(rli.condicion_ok)),
        ("cuadro", rli.cuadro),
    ]


def _escribir_vectores(ws, vectores):
    ws.append(["RUT", "Periodo", "Id", "Glosa", "Valor", "Existe", "query"])
    for codigo, valor in vectores.items():
        ws.append([RUT_CASO, 2026, codigo, GLOSAS.get(codigo, codigo), int(valor), "S", None])


def _escribir_calculadora(ws):
    ws.append(["RUT", "Periodo", "Id", "Valor", "Existe", "Fecha", "Origen", "query"])


def _escribir_retiros(ws, filas):
    ws.append([
        "RUT", "Usufructuario", "Acciones", "F1_Fecha", "F1_Monto", "F1_ISFUT_H",
        "F1_ISFUT_A", "Saldo", "F2_Fecha", "F2_Monto", "F2_ISFUT_H", "F2_ISFUT_A",
    ])
    for fila in filas:
        ws.append([
            fila["rut"], fila["usufructuario"], fila["acciones"], fila["f1_fecha"],
            fila["f1_monto"], fila["f1_isfut_h"], fila["f1_isfut_a"], fila["saldo"],
            fila["f2_fecha"], fila["f2_monto"], fila["f2_isfut_h"], fila["f2_isfut_a"],
        ])


def _escribir_guia(ws, caso, base, con_si, con_no):
    negrita = Font(bold=True)
    ws.append(["Caso", caso["titulo"]])
    ws.append(["Descripcion", caso["descripcion"]])
    ws.append(["Cuadro esperado (inicial)", caso["cuadro_esperado"]])
    ws.append([])
    ws.append(["Pasos en la UI"])
    for paso in caso["pasos"]:
        ws.append([paso])
    ws.append([])
    ws.append(["Resultado esperado (estado inicial)"])
    for etiqueta, valor in _tabla(base):
        ws.append([etiqueta, valor])
    ws.append([])
    ws.append(["Referencia: respuesta SI al 14E"])
    for etiqueta, valor in _tabla(con_si):
        ws.append([etiqueta, valor])
    ws.append([])
    ws.append(["Referencia: respuesta NO al 14E"])
    for etiqueta, valor in _tabla(con_no):
        ws.append([etiqueta, valor])
    ws["A1"].font = negrita


def _construir(caso):
    retiros = caso.get("retiros")
    base = _calcular(caso["vectores"], caso.get("externos"), retiros=retiros, acoge=None)
    if base.cuadro != caso["cuadro_esperado"]:
        raise SystemExit(
            "Cuadro inesperado en %s: %s != %s"
            % (caso["archivo"], base.cuadro, caso["cuadro_esperado"])
        )
    con_si = _calcular(caso["vectores"], caso.get("externos"), retiros=retiros, acoge=True)
    con_no = _calcular(caso["vectores"], caso.get("externos"), retiros=retiros, acoge=False)

    wb = Workbook()
    ws_v = wb.active
    ws_v.title = "Vectores"
    _escribir_vectores(ws_v, caso["vectores"])
    _escribir_calculadora(wb.create_sheet("Calculadora"))
    if retiros:
        _escribir_retiros(wb.create_sheet("Retiros"), retiros)
    _escribir_guia(wb.create_sheet("Guia"), caso, base, con_si, con_no)
    wb.save(SALIDA / caso["archivo"])
    return base, con_si


def main():
    SALIDA.mkdir(parents=True, exist_ok=True)
    print("Salida:", SALIDA)
    for caso in CASOS:
        base, con_si = _construir(caso)
        print(
            "OK %-42s cuadro=%-9s 9.1=%-11s 9.6=%-11s (SI->9.4=%s)"
            % (caso["archivo"], base.cuadro, base.v9_1, base.v9_6, con_si.v9_4)
        )


if __name__ == "__main__":
    main()

